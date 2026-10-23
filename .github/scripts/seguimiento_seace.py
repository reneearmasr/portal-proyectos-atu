"""
Seguimiento de procesos de selección de la ATU en el SEACE
==========================================================

Consulta los procedimientos de selección convocados por la Autoridad de
Transporte Urbano para Lima y Callao (RUC 20604932964) y genera un Excel con
el resumen de cada proceso vigente (estado, cronograma, postores, buena pro,
contrato y documentos).

Fuente: el buscador público del SEACE
(prod2.seace.gob.pe/seacebus-uiwd-pub) está protegido con reCAPTCHA, por lo
que no se consulta de forma automatizada. Se usa la API oficial de datos
abiertos del OECE (estándar OCDS), que publica los mismos procesos del SEACE
y se actualiza a diario:
    https://contratacionesabiertas.oece.gob.pe/api

Si está configurado en config.ini, envía el Excel por correo y por Telegram
(adjunto) y un resumen breve por WhatsApp.

Uso:
    py seguimiento_seace.py                    # vigentes convocados desde el año pasado
    py seguimiento_seace.py --desde-anio 2026  # solo convocatorias 2026
    py seguimiento_seace.py --todos            # incluye procesos ya culminados
    py seguimiento_seace.py --no-enviar        # solo genera el Excel
"""

import argparse
import configparser
import json
import os
import shutil
import smtplib
import sys
import time
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from email.message import EmailMessage
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE_DIR = Path(__file__).resolve().parent
API = "https://contratacionesabiertas.oece.gob.pe/api/v1"
URL_PROCESO = "https://contratacionesabiertas.oece.gob.pe/proceso/{ocid}"
# La API rechaza (403) el User-Agent por defecto de requests
HEADERS = {"User-Agent": "ATU-SeguimientoSEACE/1.0 (consulta datos abiertos OCDS)"}

ENTIDAD = "AUTORIDAD DE TRANSPORTE URBANO PARA LIMA Y CALLAO"
RUC = "20604932964"
BUYER_ID = "PE-CONSUCODE-204448"  # código de la entidad en el SEACE

# Estados de ítem en los que el procedimiento ya terminó (no está vigente)
ESTADOS_FINALES = {
    "CONTRATADO", "DESIERTO", "NULO", "CANCELADO", "CULMINADO",
    "NO_SUSCRIPCION_CONTRATO POR_DECISION_ENTIDAD",
}

# Un proceso "CONVOCADO" sin documentos nuevos por más días se marca para verificar
DIAS_SIN_MOVIMIENTO = 60

OBJETOS = {"goods": "Bienes", "services": "Servicios", "works": "Obras",
           "consultingServices": "Consultoría"}

TIPOS_DOC = {
    "biddingDocuments": "Bases", "clarifications": "Absolución de consultas",
    "evaluationReports": "Evaluación / cuadro comparativo", "awardNotice": "Buena pro",
    "contractSigned": "Contrato", "tenderNotice": "Convocatoria",
    "complaints": "Recurso / apelación", "technicalSpecifications": "Especificaciones técnicas",
}


@dataclass
class Proceso:
    ocid: str
    convocatoria: str
    nomenclatura: str
    tipo: str
    objeto: str
    descripcion: str
    fecha_convocatoria: datetime | None
    estados: list
    items: list
    valor: float
    moneda: str
    consultas: tuple
    periodo: tuple
    postores: list
    adjudicaciones: list
    contratos: list
    documentos: list
    url: str
    resumen: str = ""
    situacion: str = ""
    extra: dict = field(default_factory=dict)

    @property
    def estado(self) -> str:
        return " / ".join(self.estados) if self.estados else "SIN ESTADO"

    @property
    def vigente(self) -> bool:
        return any(e not in ESTADOS_FINALES for e in self.estados) or not self.estados


# --------------------------------------------------------------------------
# Consulta a la API OCDS
# --------------------------------------------------------------------------
def fecha(valor) -> datetime | None:
    if not valor:
        return None
    try:
        return datetime.fromisoformat(valor).replace(tzinfo=None)
    except ValueError:
        return None


def obtener_procesos(desde_anio: int, por_pagina=50):
    """Recorre los procesos de la entidad, del más reciente al más antiguo."""
    s = requests.Session()
    s.headers.update(HEADERS)
    pagina, crudos = 1, []
    while True:
        params = {"buyerID": BUYER_ID, "processesPage": pagina,
                  "processesPaginateBy": por_pagina, "format": "json",
                  "order_processes_date": "desc"}
        for intento in range(3):
            try:
                r = s.get(f"{API}/buyerProcesses", params=params, timeout=120)
                r.raise_for_status()
                break
            except requests.RequestException:
                if intento == 2:
                    raise
                time.sleep(3 * (intento + 1))
        datos = r.json()
        lote = [x["compiledRelease"] for x in datos.get("results", [])]
        crudos.extend(lote)
        ultimo = fecha(lote[-1].get("tender", {}).get("datePublished")) if lote else None
        pag = datos.get("pagination") or {}
        if not pag.get("has_next") or (ultimo and ultimo.year < desde_anio):
            break
        pagina += 1
    return [c for c in crudos
            if (f := fecha(c.get("tender", {}).get("datePublished"))) and f.year >= desde_anio]


def a_proceso(c: dict) -> Proceso:
    t = c.get("tender", {})
    items = t.get("items", [])
    estados = []
    for it in items:
        e = (it.get("statusDetails") or "").strip()
        if e and e not in estados:
            estados.append(e)

    valor = (t.get("value") or {}).get("amount") or 0.0
    if not valor:
        valor = sum((it.get("totalValue") or {}).get("amount") or 0 for it in items)
    moneda = (t.get("value") or {}).get("currency", "PEN")

    adjudicaciones = []
    for a in c.get("awards", []):
        adjudicaciones.append({
            "proveedores": ", ".join(p.get("name", "") for p in a.get("suppliers", [])),
            "monto": (a.get("value") or {}).get("amount"),
            "fecha": fecha(a.get("date")),
        })
    contratos = []
    for k in c.get("contracts", []):
        contratos.append({
            "titulo": k.get("title") or k.get("id", ""),
            "monto": (k.get("value") or {}).get("amount"),
            "firma": fecha(k.get("dateSigned")),
        })
    documentos = sorted(
        ({"titulo": d.get("title", ""), "tipo": TIPOS_DOC.get(d.get("documentType"), d.get("documentType", "")),
          "fecha": fecha(d.get("datePublished")), "url": d.get("url", "")}
         for d in t.get("documents", [])),
        key=lambda d: d["fecha"] or datetime.min)

    ep, tp = t.get("enquiryPeriod") or {}, t.get("tenderPeriod") or {}
    return Proceso(
        ocid=c["ocid"],
        convocatoria=str(t.get("id", "")),
        nomenclatura=t.get("title", ""),
        tipo=t.get("procurementMethodDetails", ""),
        objeto=OBJETOS.get(t.get("mainProcurementCategory"), t.get("mainProcurementCategory", "")),
        descripcion=(t.get("description") or "").strip().rstrip("¿").strip(),
        fecha_convocatoria=fecha(t.get("datePublished")),
        estados=estados,
        items=items,
        valor=valor,
        moneda=moneda,
        consultas=(fecha(ep.get("startDate")), fecha(ep.get("endDate"))),
        periodo=(fecha(tp.get("startDate")), fecha(tp.get("endDate"))),
        postores=[p.get("name", "") for p in t.get("tenderers", [])],
        adjudicaciones=adjudicaciones,
        contratos=contratos,
        documentos=documentos,
        url=URL_PROCESO.format(ocid=c["ocid"]),
    )


# --------------------------------------------------------------------------
# Resumen por proceso
# --------------------------------------------------------------------------
def dias_entre(ini: datetime | None, fin: datetime) -> int | None:
    """Días calendario transcurridos (por fecha, sin considerar la hora)."""
    return (fin.date() - ini.date()).days if ini else None


def soles(monto) -> str:
    return f"S/ {monto:,.2f}" if monto else "no publicado"


def describir(p: Proceso, hoy: datetime):
    """Arma el texto de resumen y la situación / próxima acción del proceso."""
    dias = dias_entre(p.fecha_convocatoria, hoy)
    tipo = p.tipo if p.objeto.lower() in p.tipo.lower() else f"{p.tipo} de {p.objeto.lower()}"
    partes = [f"{tipo} convocado el {p.fecha_convocatoria:%d/%m/%Y}"
              + (f" (hace {dias} días)" if dias else " (hoy)" if dias == 0 else "") + "."]
    partes.append(f"Valor estimado/referencial: {soles(p.valor)}.")
    if len(p.items) > 1:
        conteo = Counter(it.get("statusDetails") for it in p.items)
        partes.append(f"{len(p.items)} ítems: " + ", ".join(f"{n} {e}" for e, n in conteo.items()) + ".")
    else:
        partes.append(f"Estado: {p.estado}.")
    if p.postores:
        partes.append(f"Postores: {len(p.postores)}.")
    for a in p.adjudicaciones:
        partes.append(f"Buena pro a {a['proveedores']} por {soles(a['monto'])}.")
    for k in p.contratos:
        partes.append(f"Contrato {k['titulo']}" + (f" firmado el {k['firma']:%d/%m/%Y}" if k["firma"] else "")
                      + f" por {soles(k['monto'])}.")
    if p.documentos:
        d = p.documentos[-1]
        partes.append(f"Último documento: {d['titulo']}" + (f" ({d['fecha']:%d/%m/%Y})" if d["fecha"] else "") + ".")
    p.resumen = " ".join(partes)

    ini, fin = p.consultas
    estados = set(p.estados)
    ult_mov = max((d["fecha"] for d in p.documentos if d["fecha"]), default=p.fecha_convocatoria)
    sin_mov = dias_entre(ult_mov, hoy) or 0
    # grupo: clasificación usada en el portal para los indicadores y filtros
    if "APELADO" in estados:
        sit, grupo = "En apelación: hacer seguimiento a la resolución del recurso.", "apelacion"
    elif estados & {"RETROTRAIDO_POR_RESOLUCION", "CONVOCADO_POR_REINICIO"}:
        sit, grupo = "Retrotraído/reiniciado: revisar nueva etapa del cronograma.", "retrotraido"
    elif "CONSENTIDO" in estados:
        sit, grupo = "Buena pro consentida: pendiente suscripción de contrato.", "contrato"
    elif "ADJUDICADO" in estados:
        sit, grupo = "Buena pro otorgada: en plazo de consentimiento.", "contrato"
    elif "CONVOCADO" in estados and fin and fin >= hoy:
        sit, grupo = f"Consultas y observaciones hasta el {fin:%d/%m/%Y} ({dias_entre(hoy, fin)} días).", "curso"
    elif "CONVOCADO" in estados and sin_mov > DIAS_SIN_MOVIMIENTO:
        sit, grupo = (f"Convocado sin documentos nuevos hace {sin_mov} días: "
                      "verificar en SEACE si sigue en curso o no fue actualizado."), "verificar"
    elif "CONVOCADO" in estados:
        sit, grupo = "Convocado: etapa de consultas cerrada; en curso hacia presentación de ofertas/evaluación.", "curso"
    elif estados - ESTADOS_FINALES:
        sit = "En trámite: " + ", ".join(sorted(estados - ESTADOS_FINALES)).lower().replace("_", " ") + "."
        grupo = "curso"
    else:
        sit, grupo = "Culminado.", "culminado"
    p.situacion = sit
    p.extra["grupo"] = grupo


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
AZUL = "1F4E79"
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
COLORES_ESTADO = {
    "CONVOCADO": "DDEBF7", "APELADO": "F8CBAD", "CONSENTIDO": "E2EFDA", "ADJUDICADO": "E2EFDA",
    "RETROTRAIDO_POR_RESOLUCION": "FFF2CC", "CONVOCADO_POR_REINICIO": "FFF2CC",
}


def hoja_tabla(ws, fila_ini, encabezados, filas, anchos, nombre_tabla, envolver=()):
    for j, enc in enumerate(encabezados, 1):
        c = ws.cell(fila_ini, j, enc)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
    for i, fila in enumerate(filas, fila_ini + 1):
        for j, valor in enumerate(fila, 1):
            c = ws.cell(i, j, valor)
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=j in envolver)
            if isinstance(valor, datetime):
                c.number_format = "dd/mm/yyyy hh:mm" if (valor.hour or valor.minute) else "dd/mm/yyyy"
            elif isinstance(valor, float):
                c.number_format = '#,##0.00'
    for j, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = ancho
    if filas:
        ref = f"A{fila_ini}:{get_column_letter(len(encabezados))}{fila_ini + len(filas)}"
        tabla = Table(displayName=nombre_tabla, ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tabla)
    ws.freeze_panes = ws.cell(fila_ini + 1, 3)


def generar_excel(procesos, total_revisados, desde_anio, solo_vigentes, hoy, ruta: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"Procesos de selección {'vigentes ' if solo_vigentes else ''}– {ENTIDAD} (RUC {RUC})"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = (f"Convocatorias desde {desde_anio} · Revisados: {total_revisados} · "
                f"{'Vigentes' if solo_vigentes else 'Listados'}: {len(procesos)} · "
                f"Generado: {hoy:%d/%m/%Y %H:%M} · Fuente: OECE – Contrataciones Abiertas (datos SEACE)")
    ws["A2"].font = Font(italic=True, color="595959")

    enc = ["N°", "Nomenclatura", "N° convocatoria", "Tipo de procedimiento", "Objeto", "Descripción", "Fecha convocatoria",
           "Días desde convocatoria", "Estado actual", "N° ítems", "Valor estimado / referencial (S/)",
           "Consultas y observaciones (inicio)", "Consultas y observaciones (fin)", "N° postores",
           "Buena pro (proveedor)", "Monto adjudicado (S/)", "Contrato", "Último documento",
           "Fecha último documento", "Situación / próxima acción", "Resumen del proceso", "Enlace"]
    filas = []
    for n, p in enumerate(procesos, 1):
        ult = p.documentos[-1] if p.documentos else {}
        filas.append([
            n, p.nomenclatura, p.convocatoria, p.tipo, p.objeto, p.descripcion, p.fecha_convocatoria,
            dias_entre(p.fecha_convocatoria, hoy),
            p.estado, len(p.items), float(p.valor) if p.valor else None,
            p.consultas[0], p.consultas[1], len(p.postores) or None,
            "; ".join(a["proveedores"] for a in p.adjudicaciones) or None,
            float(sum(a["monto"] or 0 for a in p.adjudicaciones)) or None,
            "; ".join(k["titulo"] for k in p.contratos) or None,
            ult.get("titulo"), ult.get("fecha"), p.situacion, p.resumen, p.url,
        ])
    anchos = [5, 26, 12, 22, 11, 50, 16, 11, 22, 7, 16, 16, 16, 9, 30, 16, 20, 28, 16, 40, 70, 18]
    hoja_tabla(ws, 4, enc, filas, anchos, "Procesos", envolver={6, 15, 18, 20, 21})
    for i, p in enumerate(procesos, 5):
        color = next((COLORES_ESTADO[e] for e in p.estados if e in COLORES_ESTADO), None)
        if color:
            ws.cell(i, 9).fill = PatternFill("solid", fgColor=color)
        c = ws.cell(i, 22)
        c.value, c.hyperlink, c.font = "Ver proceso", p.url, Font(color="0563C1", underline="single")
    ws.row_dimensions[4].height = 45
    if not procesos:
        ws["A5"] = "No se encontraron procesos vigentes en el periodo."

    # Detalle por ítem
    wi = wb.create_sheet("Ítems")
    filas = []
    for p in procesos:
        for it in p.items:
            filas.append([p.nomenclatura, it.get("position"), it.get("description", "").strip().rstrip("¿"),
                          it.get("statusDetails"), it.get("quantity"), (it.get("unit") or {}).get("name"),
                          float((it.get("totalValue") or {}).get("amount") or 0) or None,
                          (it.get("classification") or {}).get("description")])
    hoja_tabla(wi, 1, ["Nomenclatura", "Ítem", "Descripción", "Estado", "Cantidad", "Unidad",
                       "Valor (S/)", "Clasificación CUBSO"], filas,
               [26, 6, 60, 22, 10, 12, 16, 35], "Items", envolver={3, 8})

    # Documentos publicados
    wd = wb.create_sheet("Documentos")
    filas = [[p.nomenclatura, d["fecha"], d["tipo"], d["titulo"], d["url"]]
             for p in procesos for d in p.documentos]
    hoja_tabla(wd, 1, ["Nomenclatura", "Fecha publicación", "Tipo", "Documento", "Descarga"], filas,
               [26, 18, 26, 45, 14], "Documentos", envolver={4})
    for i, fila in enumerate(filas, 2):
        if fila[4]:
            c = wd.cell(i, 5)
            c.value, c.hyperlink, c.font = "Descargar", fila[4], Font(color="0563C1", underline="single")

    # Conteo por estado
    we = wb.create_sheet("Por estado")
    conteo = Counter(p.estado for p in procesos)
    hoja_tabla(we, 1, ["Estado", "N° procesos"], sorted(conteo.items(), key=lambda x: -x[1]),
               [40, 12], "PorEstado")

    wb.save(ruta)


# --------------------------------------------------------------------------
# Archivos para el Portal de Proyectos ATU (GitHub Pages)
# --------------------------------------------------------------------------
NOMBRE_EXCEL_PORTAL = "Procesos_SEACE_ATU.xlsx"


def exportar_portal(procesos, total, desde_anio, hoy, ruta_xlsx: Path, carpeta: Path):
    """Deja en la carpeta del portal procesos.json y el Excel con nombre fijo."""
    carpeta.mkdir(parents=True, exist_ok=True)
    iso = lambda f: f.isoformat(timespec="minutes") if f else None  # noqa: E731
    datos = {
        "generado": hoy.isoformat(timespec="minutes"),
        "entidad": ENTIDAD, "ruc": RUC, "desde_anio": desde_anio, "revisados": total,
        "fuente": "OECE – Portal de Contrataciones Abiertas (datos del SEACE)",
        "excel": NOMBRE_EXCEL_PORTAL,
        "procesos": [{
            "nomenclatura": p.nomenclatura, "convocatoria": p.convocatoria, "tipo": p.tipo,
            "objeto": p.objeto, "descripcion": p.descripcion,
            "fecha_convocatoria": iso(p.fecha_convocatoria),
            "estado": p.estado, "grupo": p.extra.get("grupo", ""),
            "items": len(p.items), "valor": p.valor or None,
            "consultas_fin": iso(p.consultas[1]), "postores": len(p.postores),
            "buena_pro": "; ".join(a["proveedores"] for a in p.adjudicaciones),
            "monto_adjudicado": sum(a["monto"] or 0 for a in p.adjudicaciones) or None,
            "ultimo_documento": p.documentos[-1]["titulo"] if p.documentos else "",
            "fecha_ultimo_documento": iso(p.documentos[-1]["fecha"]) if p.documentos else None,
            "situacion": p.situacion, "resumen": p.resumen, "url": p.url,
        } for p in procesos],
    }
    (carpeta / "procesos.json").write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    shutil.copyfile(ruta_xlsx, carpeta / NOMBRE_EXCEL_PORTAL)
    print(f"  Portal actualizado: {carpeta} (suba procesos.json y {NOMBRE_EXCEL_PORTAL} a GitHub)")


# --------------------------------------------------------------------------
# Envío
# --------------------------------------------------------------------------
def texto_breve(procesos, total, hoy):
    lineas = [f"*Procesos SEACE vigentes – ATU ({hoy:%d/%m/%Y})*",
              f"Revisados: {total} | Vigentes: {len(procesos)}"]
    for p in procesos:
        lineas.append(f"\n*{p.nomenclatura}* – {p.estado}\n{p.descripcion[:110]}"
                      f"{'…' if len(p.descripcion) > 110 else ''}\n{p.situacion}")
    if not procesos:
        lineas.append("No hay procesos vigentes.")
    return "\n".join(lineas)


def enviar_correo(cfg, ruta_xlsx: Path, cuerpo: str, asunto: str):
    c = cfg["correo"]
    clave = os.environ.get("SEACE_SMTP_PASSWORD") or c.get("password", "")
    if not (c.get("usuario") and clave and c.get("destinatarios")):
        print("[correo] Falta usuario, password o destinatarios en config.ini; no se envía.")
        return
    msg = EmailMessage()
    msg["Subject"] = asunto
    msg["From"] = c.get("usuario")
    msg["To"] = c.get("destinatarios")
    msg.set_content(cuerpo.replace("*", "") + "\n\nSe adjunta el resumen en Excel.")
    msg.add_attachment(ruta_xlsx.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       filename=ruta_xlsx.name)
    with smtplib.SMTP(c.get("servidor", "smtp.gmail.com"), c.getint("puerto", 587), timeout=60) as srv:
        srv.starttls()
        srv.login(c.get("usuario"), clave)
        srv.send_message(msg)
    print(f"[correo] Enviado a {c.get('destinatarios')}")


def enviar_whatsapp(cfg, cuerpo: str):
    """Envía un texto vía CallMeBot (https://www.callmebot.com/blog/free-api-whatsapp-messages/)."""
    w = cfg["whatsapp"]
    telefono, apikey = w.get("telefono", ""), w.get("apikey", "")
    if not (telefono and apikey):
        print("[whatsapp] Falta telefono o apikey en config.ini; no se envía.")
        return
    # CallMeBot admite mensajes cortos: se parte en bloques
    bloques, actual = [], ""
    for linea in cuerpo.split("\n"):
        if len(actual) + len(linea) > 900:
            bloques.append(actual); actual = ""
        actual += linea + "\n"
    bloques.append(actual)
    for b in bloques:
        r = requests.get("https://api.callmebot.com/whatsapp.php",
                         params={"phone": telefono, "text": b, "apikey": apikey}, timeout=60)
        if r.status_code != 200:
            print(f"[whatsapp] Error {r.status_code}: {r.text[:200]}")
            return
        time.sleep(3)
    print(f"[whatsapp] Enviado ({len(bloques)} mensaje/s)")


def enviar_telegram(cfg, ruta_xlsx: Path, cuerpo: str):
    """Envía el resumen y el Excel por un bot de Telegram (creado con @BotFather)."""
    t = cfg["telegram"]
    token = os.environ.get("SEACE_TELEGRAM_TOKEN") or t.get("token", "")
    chat_id = os.environ.get("SEACE_TELEGRAM_CHAT_ID") or t.get("chat_id", "")
    if not token:
        print("[telegram] Falta token en config.ini o SEACE_TELEGRAM_TOKEN; no se envía.")
        return
    api = f"https://api.telegram.org/bot{token}"
    if not chat_id:
        # El chat_id aparece después de escribirle /start al bot
        r = requests.get(f"{api}/getUpdates", timeout=60).json()
        chats = {u["message"]["chat"]["id"] for u in r.get("result", []) if "message" in u}
        if not chats:
            print("[telegram] Escriba /start a su bot en Telegram y vuelva a ejecutar.")
            return
        chat_id = str(chats.pop())
        print(f"[telegram] chat_id detectado: {chat_id} (puede fijarlo en config.ini)")

    # Telegram limita los mensajes a 4096 caracteres
    bloques, actual = [], ""
    for linea in cuerpo.replace("*", "").split("\n"):
        if len(actual) + len(linea) > 3800:
            bloques.append(actual); actual = ""
        actual += linea + "\n"
    bloques.append(actual)
    for b in bloques:
        r = requests.post(f"{api}/sendMessage", data={"chat_id": chat_id, "text": b}, timeout=60)
        if not r.ok:
            print(f"[telegram] Error {r.status_code}: {r.text[:200]}")
            return
    with ruta_xlsx.open("rb") as f:
        r = requests.post(f"{api}/sendDocument", data={"chat_id": chat_id},
                          files={"document": (ruta_xlsx.name, f)}, timeout=120)
    if not r.ok:
        print(f"[telegram] Error al adjuntar Excel {r.status_code}: {r.text[:200]}")
        return
    print(f"[telegram] Enviado resumen y Excel ({len(bloques)} mensaje/s)")


# --------------------------------------------------------------------------
def main():
    hoy = datetime.now()
    ap = argparse.ArgumentParser(description="Seguimiento de procesos SEACE de la ATU")
    ap.add_argument("--desde-anio", type=int, default=hoy.year - 1,
                    help="año mínimo de convocatoria (por defecto: año pasado)")
    ap.add_argument("--todos", action="store_true", help="incluir procesos culminados")
    ap.add_argument("--no-enviar", action="store_true", help="solo genera el Excel")
    ap.add_argument("--portal", help="carpeta donde dejar procesos.json y el Excel del portal "
                                     "(por defecto: carpeta_portal de config.ini)")
    args = ap.parse_args()

    cfg = configparser.ConfigParser()
    cfg.read(BASE_DIR / "config.ini", encoding="utf-8")
    for sec in ("correo", "whatsapp", "telegram", "general"):
        if not cfg.has_section(sec):
            cfg.add_section(sec)

    print(f"Consultando procesos de la ATU convocados desde {args.desde_anio}…")
    procesos = [a_proceso(c) for c in obtener_procesos(args.desde_anio)]
    total = len(procesos)
    for p in procesos:
        describir(p, hoy)
    if not args.todos:
        procesos = [p for p in procesos if p.vigente]
    print(f"  {total} procesos revisados; {len(procesos)} {'listados' if args.todos else 'vigentes'}.")

    carpeta = Path(cfg["general"].get("carpeta_salida", BASE_DIR / "resumenes"))
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / f"Procesos_SEACE_ATU_{hoy:%Y%m%d_%H%M}.xlsx"
    generar_excel(procesos, total, args.desde_anio, not args.todos, hoy, ruta)
    print(f"  Excel generado: {ruta}")

    if not args.todos:
        portal = Path(args.portal or cfg["general"].get(
            "carpeta_portal", BASE_DIR.parent / "portal-proyectos-atu" / "seace"))
        if portal.parent.exists():
            exportar_portal(procesos, total, args.desde_anio, hoy, ruta, portal)

    if args.no_enviar:
        return
    cuerpo = texto_breve(procesos, total, hoy)
    asunto = f"Procesos SEACE ATU – {len(procesos)} vigentes ({hoy:%d/%m/%Y})"
    try:
        enviar_correo(cfg, ruta, cuerpo, asunto)
    except Exception as e:  # noqa: BLE001
        print(f"[correo] Error: {e}")
    try:
        enviar_whatsapp(cfg, cuerpo)
    except Exception as e:  # noqa: BLE001
        print(f"[whatsapp] Error: {e}")
    try:
        enviar_telegram(cfg, ruta, cuerpo)
    except Exception as e:  # noqa: BLE001
        token = os.environ.get("SEACE_TELEGRAM_TOKEN") or cfg["telegram"].get("token", "")
        print(f"[telegram] Error: {str(e).replace(token, '<token>') if token else e}")
        if "CERTIFICATE_VERIFY_FAILED" in str(e):
            print("[telegram] La red intercepta HTTPS (p. ej. firewall FortiGate). Ejecute desde otra "
                  "red o pida a TI instalar el certificado raíz del firewall en esta PC.")


if __name__ == "__main__":
    sys.exit(main())
